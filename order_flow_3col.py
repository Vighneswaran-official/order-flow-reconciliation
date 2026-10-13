#!/usr/bin/env python3
"""
Order Flow & Seller SKU Lookup + Status Mapping Automation (3-Column Output)
============================================================================
Produces an exact 3-column enriched dataset from two CSV reports:
  Column 1: item_id
  Column 2: seller sku code
  Column 3: final_status

INPUT FILES (auto-detected from ./input folder):
  1. Seller Orders Report -> filename contains "Seller_Orders_Report"
     - Key column:   "order line id"
     - Fetch column: "seller sku code"
  2. Order Flow Report   -> filename starts with "part-"
     - Key column:   "item_id"
     - Also uses:    "return_type"

LOGIC:
  - All columns read strictly as text (pandas dtype=str, keep_default_na=False).
    IDs never convert to floats or scientific notation.
  - Strips whitespace & leading apostrophes (') from key columns.
  - Base table = Order Flow report (keeps every single row).
  - Left join: Order Flow "item_id" == Orders "order line id".
    * Unmatched -> "NOT FOUND"
    * Duplicate keys in Orders -> retains first occurrence and logs warning.
  - Case-insensitive mapping of "return_type" using the dictionary below.
    * blank / empty / NaN -> "Delivered"
    * release_rto_refund  -> "RTO"
    * return_refund       -> "RVP"
    * exchange            -> "Exchange"
    * unknown             -> "UNMAPPED: <raw value>" and logs warning.
  - Outputs 3-column file to ./output as both CSV and Excel (.xlsx) with strict text formatting.
"""

import os
import sys
import glob
import argparse
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==============================================================================
# CONFIGURATION: RETURN TYPE TO FINAL STATUS MAPPING
# Edit or add new mappings here easily
# ==============================================================================
RETURN_TYPE_STATUS_MAPPING = {
    "release_rto_refund": "RTO",
    "return_refund": "RVP",
    "exchange": "Exchange",
}
DEFAULT_BLANK_STATUS = "Delivered"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_INPUT_DIR = os.path.join(BASE_DIR, "input")
DEFAULT_OUTPUT_DIR = os.path.join(BASE_DIR, "output")
DEFAULT_LOG_PATH = os.path.join(BASE_DIR, "logs", "log_3col.txt")


class DualLogger:
    """Logs messages simultaneously to the console and to log file."""
    def __init__(self, log_path: str):
        self.log_path = log_path
        self.entries = []
        os.makedirs(os.path.dirname(os.path.abspath(log_path)), exist_ok=True)
        with open(self.log_path, "w", encoding="utf-8") as f:
            f.write(f"=== ORDER FLOW 3-COLUMN ENRICHMENT AUDIT LOG: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} ===\n\n")

    def log(self, message: str = ""):
        print(message)
        self.entries.append(message)
        with open(self.log_path, "a", encoding="utf-8") as f:
            f.write(message + "\n")

    def warning(self, message: str):
        self.log(f"[WARNING] {message}")

    def error(self, message: str):
        self.log(f"[ERROR] {message}")

    def info(self, message: str):
        self.log(message)


def clean_key_series(series: pd.Series) -> pd.Series:
    """
    Sanitizes key columns:
    - Casts to string
    - Strips leading & trailing whitespace
    - Strips leading apostrophe (') commonly exported by Excel
    """
    return (
        series.fillna("")
        .astype(str)
        .str.strip()
        .str.lstrip("'")
        .str.strip()
    )


def map_return_type_to_status(val: str, mapping: dict = None, blank_status: str = None) -> str:
    """
    Maps return_type to final_status case-insensitively.
    """
    if mapping is None:
        mapping = RETURN_TYPE_STATUS_MAPPING
    if blank_status is None:
        blank_status = DEFAULT_BLANK_STATUS

    raw = str(val).strip() if val is not None else ""
    if not raw or raw.lower() in ("nan", "none", "null"):
        return blank_status

    normalized = raw.lower()
    for map_key, map_status in mapping.items():
        if normalized == map_key.strip().lower():
            return map_status

    return f"UNMAPPED: {raw}"


def find_input_files(input_dir: str, logger: DualLogger):
    """
    Auto-detects the newest Seller Orders Report and Order Flow Report in input_dir.
    """
    if not os.path.exists(input_dir):
        raise FileNotFoundError(f"Input directory does not exist: {input_dir}")

    csv_files = glob.glob(os.path.join(input_dir, "*.csv"))
    if not csv_files:
        raise FileNotFoundError(f"No CSV files found in input directory: {input_dir}")

    orders_candidates = []
    flow_candidates = []

    for f in csv_files:
        fname = os.path.basename(f)
        if "seller_orders_report" in fname.lower():
            orders_candidates.append(f)
        elif fname.lower().startswith("part-") or "order_flow" in fname.lower():
            flow_candidates.append(f)

    if not orders_candidates:
        raise FileNotFoundError(
            f"Could not auto-detect Seller Orders Report in '{input_dir}'. "
            f"Expected a CSV filename containing 'Seller_Orders_Report'."
        )

    if not flow_candidates:
        raise FileNotFoundError(
            f"Could not auto-detect Order Flow Report in '{input_dir}'. "
            f"Expected a CSV filename starting with 'part-'."
        )

    orders_candidates.sort(key=lambda p: os.path.getmtime(p), reverse=True)
    flow_candidates.sort(key=lambda p: os.path.getmtime(p), reverse=True)

    if len(orders_candidates) > 1:
        logger.warning(
            f"Multiple Seller Orders Reports found: {[os.path.basename(x) for x in orders_candidates]}.\n"
            f"  -> Selecting newest file: '{os.path.basename(orders_candidates[0])}'"
        )
    if len(flow_candidates) > 1:
        logger.warning(
            f"Multiple Order Flow Reports found: {[os.path.basename(x) for x in flow_candidates]}.\n"
            f"  -> Selecting newest file: '{os.path.basename(flow_candidates[0])}'"
        )

    return orders_candidates[0], flow_candidates[0]


def load_csv_as_text(filepath: str, label: str, logger: DualLogger) -> pd.DataFrame:
    """
    Reads a CSV strictly as string with multiple encoding fallbacks.
    """
    encodings = ["utf-8-sig", "utf-8", "cp1252", "latin1"]
    df = None
    last_err = None

    for enc in encodings:
        try:
            df = pd.read_csv(
                filepath,
                dtype=str,
                keep_default_na=False,
                encoding=enc,
                low_memory=False
            )
            logger.info(f"Loaded {label} from '{os.path.basename(filepath)}' using encoding '{enc}' ({len(df):,} rows).")
            break
        except UnicodeDecodeError as e:
            last_err = e
            continue
        except Exception as e:
            raise RuntimeError(f"Error reading {label} file '{filepath}': {str(e)}")

    if df is None:
        raise RuntimeError(f"Failed to read {label} file '{filepath}'. Encoding error: {last_err}")

    df.columns = [str(c).strip() for c in df.columns]
    return df


def write_3col_excel(xlsx_path: str, df_3col: pd.DataFrame, logger: DualLogger):
    """
    Writes the 3-column DataFrame to an Excel workbook with strict string cell typing ('s', '@').
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Order_Flow_3Col"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="000000")
    
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")

    cols = list(df_3col.columns)

    # Header Row
    for col_idx, col_name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border
        cell.number_format = "@"

    # Data Rows (All typed as string 's' and format '@')
    for row_idx, row_values in enumerate(df_3col.itertuples(index=False), start=2):
        for col_idx, val in enumerate(row_values, start=1):
            val_str = "" if val is None else str(val)
            cell = ws.cell(row=row_idx, column=col_idx, value=val_str)
            cell.data_type = "s"
            cell.number_format = "@"
            cell.font = data_font
            cell.alignment = align_left
            cell.border = thin_border

            # Highlight NOT FOUND or UNMAPPED
            if val_str == "NOT FOUND":
                cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                cell.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
            elif val_str.startswith("UNMAPPED:"):
                cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                cell.font = Font(name="Calibri", size=10, bold=True, color="92400E")

    # Auto-adjust column widths
    for col_idx, col_name in enumerate(cols, start=1):
        max_len = len(str(col_name))
        for row in range(2, min(ws.max_row + 1, 200)):
            val_c = ws.cell(row=row, column=col_idx).value
            if val_c:
                max_len = max(max_len, len(str(val_c)))
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    wb.save(xlsx_path)


def process_3col_pipeline(
    orders_path: str = None,
    flow_path: str = None,
    input_dir: str = DEFAULT_INPUT_DIR,
    output_dir: str = DEFAULT_OUTPUT_DIR,
    log_path: str = DEFAULT_LOG_PATH
):
    """
    Executes the 3-column enrichment pipeline.
    """
    logger = DualLogger(log_path)
    logger.info("=" * 75)
    logger.info("ORDER FLOW 3-COLUMN ENRICHMENT AUTOMATION")
    logger.info("Columns Produced: ['item_id', 'seller sku code', 'final_status']")
    logger.info("=" * 75)
    logger.info(f"Timestamp        : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    # 1. Resolve files
    if not orders_path or not flow_path:
        detected_orders, detected_flow = find_input_files(input_dir, logger)
        orders_path = orders_path or detected_orders
        flow_path = flow_path or detected_flow

    logger.info(f"Seller Orders File: {orders_path}")
    logger.info(f"Order Flow File   : {flow_path}")
    logger.info(f"Output Directory  : {output_dir}")
    logger.info("-" * 75)

    # 2. Read both files strictly as text
    orders_df = load_csv_as_text(orders_path, "Seller Orders Report", logger)
    flow_df = load_csv_as_text(flow_path, "Order Flow Report", logger)

    orders_total_rows = len(orders_df)
    flow_total_rows = len(flow_df)

    # Validate columns
    if "order line id" not in orders_df.columns:
        raise KeyError(f"Missing required key 'order line id' in Seller Orders ({orders_path})")
    if "seller sku code" not in orders_df.columns:
        raise KeyError(f"Missing required column 'seller sku code' in Seller Orders ({orders_path})")
    if "item_id" not in flow_df.columns:
        raise KeyError(f"Missing required key 'item_id' in Order Flow ({flow_path})")
    if "return_type" not in flow_df.columns:
        raise KeyError(f"Missing required column 'return_type' in Order Flow ({flow_path})")

    # 3. Clean keys (trim whitespace + strip leading apostrophe)
    orders_df["_clean_order_line_id"] = clean_key_series(orders_df["order line id"])
    flow_df["_clean_item_id"] = clean_key_series(flow_df["item_id"])

    # 4. Handle duplicates in Seller Orders (keep first, log warning)
    dup_mask = orders_df.duplicated(subset=["_clean_order_line_id"], keep=False)
    dup_count = dup_mask.sum()
    if dup_count > 0:
        unique_dups = orders_df.loc[dup_mask, "_clean_order_line_id"].nunique()
        logger.warning(
            f"Seller Orders report contains {dup_count} duplicate rows ({unique_dups} unique 'order line id' keys). "
            f"Keeping the first occurrence and ignoring subsequent duplicates."
        )
        sample_dups = orders_df.loc[dup_mask, "_clean_order_line_id"].unique()[:5]
        logger.warning(f"  -> Sample duplicate keys: {list(sample_dups)}")

    orders_dedup = orders_df.drop_duplicates(subset=["_clean_order_line_id"], keep="first")

    # 5. Left join: Order Flow 'item_id' == Orders 'order line id'
    flow_df["_row_order"] = range(len(flow_df))
    orders_lookup = orders_dedup[["_clean_order_line_id", "seller sku code"]].copy()
    orders_lookup.rename(columns={"seller sku code": "_fetched_seller_sku"}, inplace=True)

    merged_df = flow_df.merge(
        orders_lookup,
        left_on="_clean_item_id",
        right_on="_clean_order_line_id",
        how="left"
    )
    merged_df.sort_values(by="_row_order", inplace=True)

    # Fill unmatched with 'NOT FOUND' (do not drop row)
    merged_df["seller sku code"] = merged_df["_fetched_seller_sku"].fillna("NOT FOUND")
    merged_df.loc[merged_df["seller sku code"].str.strip() == "", "seller sku code"] = "NOT FOUND"

    # 6. Map return_type -> final_status
    merged_df["final_status"] = merged_df["return_type"].apply(
        lambda v: map_return_type_to_status(v, mapping=RETURN_TYPE_STATUS_MAPPING, blank_status=DEFAULT_BLANK_STATUS)
    )

    unmapped_mask = merged_df["final_status"].str.startswith("UNMAPPED: ")
    unmapped_count = int(unmapped_mask.sum())
    if unmapped_count > 0:
        unmapped_vals = merged_df.loc[unmapped_mask, "return_type"].unique().tolist()
        logger.warning(f"Found {unmapped_count} row(s) with UNMAPPED 'return_type' values: {unmapped_vals}")

    # 7. Construct EXACT 3-Column DataFrame
    three_col_df = pd.DataFrame({
        "item_id": merged_df["item_id"],
        "seller sku code": merged_df["seller sku code"],
        "final_status": merged_df["final_status"]
    })

    # Validation stats
    matched_count = int((three_col_df["seller sku code"] != "NOT FOUND").sum())
    not_found_count = flow_total_rows - matched_count
    match_rate = (matched_count / flow_total_rows * 100) if flow_total_rows > 0 else 0.0

    logger.info("-" * 75)
    logger.info("AUDIT SUMMARY:")
    logger.info(f"  * Total Order Flow Rows: {flow_total_rows}")
    logger.info(f"  * Total Orders Rows    : {orders_total_rows}")
    logger.info(f"  * Matched Rows         : {matched_count} ({match_rate:.2f}%)")
    logger.info(f"  * NOT FOUND Rows       : {not_found_count}")
    logger.info(f"  * Duplicate Keys       : {dup_count} handled")
    logger.info(f"  * Unmapped Return Types: {unmapped_count}")
    logger.info("  * Final Status Breakdown:")
    for status, cnt in three_col_df["final_status"].value_counts().items():
        pct = (cnt / flow_total_rows * 100)
        logger.info(f"      - {status}: {cnt} ({pct:.2f}%)")
    logger.info("-" * 75)

    # 8. Save Output Files
    os.makedirs(output_dir, exist_ok=True)
    today_str = datetime.now().strftime("%Y-%m-%d")

    csv_out_path = os.path.join(output_dir, f"Order_Flow_3Col_{today_str}.csv")
    xlsx_out_path = os.path.join(output_dir, f"Order_Flow_3Col_{today_str}.xlsx")

    # CSV Export
    three_col_df.to_csv(csv_out_path, index=False, encoding="utf-8-sig")
    logger.info(f"Saved 3-Column CSV to   : {csv_out_path}")

    # Excel Export with text cell formatting
    write_3col_excel(xlsx_out_path, three_col_df, logger)
    logger.info(f"Saved 3-Column Excel to : {xlsx_out_path}")

    logger.info("=" * 75)
    logger.info("3-COLUMN ENRICHMENT COMPLETED SUCCESSFULLY.")
    logger.info("=" * 75)

    return {
        "csv_path": csv_out_path,
        "xlsx_path": xlsx_out_path,
        "total_rows": flow_total_rows,
        "matched_count": matched_count,
        "not_found_count": not_found_count,
        "match_rate": round(match_rate, 2),
        "three_col_df": three_col_df
    }


def main():
    parser = argparse.ArgumentParser(description="Order Flow & Seller SKU Lookup (3-Column Output)")
    parser.add_argument("--orders-file", type=str, default=None, help="Explicit path to Seller Orders CSV")
    parser.add_argument("--flow-file", type=str, default=None, help="Explicit path to Order Flow CSV")
    parser.add_argument("--input-dir", type=str, default=DEFAULT_INPUT_DIR, help="Input directory path")
    parser.add_argument("--output-dir", type=str, default=DEFAULT_OUTPUT_DIR, help="Output directory path")
    args = parser.parse_args()

    process_3col_pipeline(
        orders_path=args.orders_file,
        flow_path=args.flow_file,
        input_dir=args.input_dir,
        output_dir=args.output_dir
    )


if __name__ == "__main__":
    main()
